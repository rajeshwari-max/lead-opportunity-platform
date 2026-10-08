"""Extract reviewed hierarchy labels from the completed Excel review queue.

The workbook remains the human-readable review artifact.  This command creates
the compact JSONL file used by training and by the optional database import.
Rows marked ``Needs discussion`` are retained with that status so training can
exclude them explicitly instead of silently treating their provisional labels
as approved.
"""
from __future__ import annotations

import argparse
import json
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook


REQUIRED = {
    "Record ID", "Title", "Human status", "Reviewed category",
    "Reviewed brands (semicolon-separated)",
    "Reviewed archetypes (semicolon-separated)",
    "Reviewed Devsol verticals (semicolon-separated)",
}


def _labels(value) -> list[str]:
    return [part.strip() for part in str(value or "").split(";") if part.strip()]


def extract(workbook_path: Path) -> list[dict]:
    workbook = load_workbook(workbook_path, read_only=True, data_only=False)
    if "Review Queue" not in workbook.sheetnames:
        raise ValueError("The workbook has no 'Review Queue' worksheet.")
    sheet = workbook["Review Queue"]
    headers = [cell.value for cell in sheet[4]]
    positions = {str(value or "").strip(): index for index, value in enumerate(headers)}
    missing = REQUIRED - set(positions)
    if missing:
        raise ValueError(f"Review Queue is missing columns: {', '.join(sorted(missing))}")

    rows: list[dict] = []
    for values in sheet.iter_rows(min_row=5, values_only=True):
        record_id = str(values[positions["Record ID"]] or "").strip()
        if not record_id:
            continue
        status = str(values[positions["Human status"]] or "").strip()
        rows.append({
            "record_id": record_id,
            "title": str(values[positions["Title"]] or "").strip(),
            "human_status": status,
            "category": str(values[positions["Reviewed category"]] or "").strip(),
            "brands": _labels(values[positions["Reviewed brands (semicolon-separated)"]]),
            "archetypes": _labels(values[positions["Reviewed archetypes (semicolon-separated)"]]),
            "devsol_verticals": _labels(values[positions["Reviewed Devsol verticals (semicolon-separated)"]]),
        })
    return rows


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--out", type=Path, required=True)
    args = parser.parse_args()

    rows = extract(args.workbook)
    args.out.parent.mkdir(parents=True, exist_ok=True)
    with args.out.open("w", encoding="utf-8") as handle:
        for row in rows:
            handle.write(json.dumps(row, ensure_ascii=False, sort_keys=True) + "\n")

    statuses = Counter(row["human_status"] or "(blank)" for row in rows)
    print(f"Wrote {len(rows)} review rows to {args.out}")
    for status, count in statuses.most_common():
        print(f"  {status}: {count}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())


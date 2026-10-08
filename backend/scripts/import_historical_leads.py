"""Normalize and idempotently import historical opportunity outcomes.

Accepts the supplied OST Excel export or the versioned JSONL produced with
``--export-jsonl``.  Exact repeated titles are collapsed before import so one
lead cannot inflate outcome statistics or evaluation metrics.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import re
import sys
from datetime import date, datetime, timezone
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from app.database.db import SessionLocal, init_db  # noqa: E402
from app.database.models import HistoricalLead, ModelRegistry  # noqa: E402
from app.services.ml_hierarchy import classify_hierarchy_batch  # noqa: E402

EMAIL = re.compile(r"\b[A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,}\b", re.I)
SPACES = re.compile(r"\s+")
ALLOWED_ITEM_TYPES = {"task", "rm ask task"}
ROW_FIELDS = (
    "external_id", "title", "source", "opportunity_url", "opportunity_type",
    "sector", "verticals", "brands", "archetypes", "geography", "focus_area",
    "eligibility_text", "funding", "company_relevance", "action_taken", "status",
    "outcome", "won_lost", "reason", "reviewer_comment", "event_date",
)
RAW_DATA_FIELDS = ("status_group", "project_or_folder", "start_date", "due_date")


def text(value) -> str:
    return SPACES.sub(" ", str(value or "")).strip()


def safe_title(value) -> str:
    return EMAIL.sub("[email removed]", text(value))[:4000]


def safe_text(value, *, limit: int = 4000) -> str:
    """Normalize free text and remove addresses from the historical dataset."""
    return EMAIL.sub("[email removed]", text(value))[:limit]


def fingerprint_for(title: str) -> str:
    """Stable identity used for exact-title deduplication and repeat-safe imports."""
    return hashlib.sha256(text(title).casefold().encode("utf-8")).hexdigest()


def iso_date(value) -> str | None:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return None


def status_mapping(status: str) -> tuple[str, str, str, str]:
    lower = status.casefold()
    reason = ""
    mappings = (
        ("not strategic", "Outside strategic focus"),
        ("no time", "Insufficient internal time"),
        ("no partner", "No suitable partner"),
        ("limited bud", "Funding or budget was insufficient"),
        ("difficult geo", "Geography was unsuitable"),
        ("bad past exp", "Previous experience was unfavorable"),
        ("missed deadline", "Deadline was missed"),
        ("low feasibility", "Feasibility was too low"),
        ("no capacity", "Insufficient internal capacity"),
        ("client cancelled", "Opportunity was cancelled"),
        ("conf.of interest", "Conflict of interest"),
        ("not relevant", "Not relevant to the company"),
        ("high risk", "Risk was too high"),
        ("forgot", "Submission was not completed"),
    )
    for needle, explanation in mappings:
        if needle in lower:
            reason = explanation
            break
    if "contracted" in lower:
        return "Won", "won", "Accepted", reason or "Contracted"
    # Wrike merged three materially different outcomes into this one bucket:
    # an award, an accepted invitation and simple event attendance. Preserve it
    # as a useful positive signal, but never claim it is a verified commercial
    # win. The same rule applies to alliance/platform acceptance.
    if any(term in lower for term in (
        "awarded", "accepted", "attended", "appd to join", "platform joined",
    )):
        return "Positive / accepted", "", "Accepted", reason or "Positive outcome; exact result was not distinguished"
    # The legacy bucket explicitly contains lost/rejected. It is the only
    # defensible negative won/lost label in this export. Not-pursued,
    # not-submitted and cancelled records remain rejection evidence without
    # being misrepresented as bids the company actually lost.
    if "lost/rejected" in lower or lower.startswith("lost"):
        return "Lost or rejected", "lost", "Rejected", reason or "Lost or rejected"
    if (
        lower.startswith("12.")
        or any(term in lower for term in (
            "don't pursue", "dont pursue", "not sub", "not submitted",
            "cancelled", "not attend", "not relevant",
        ))
    ):
        return "Not pursued / not submitted", "", "Rejected", reason or "Not pursued or not submitted"
    if "shortlist" in lower:
        return "Shortlisted", "", "Shortlisted", reason
    if "submitted" in lower or "waiting" in lower:
        return "Submitted", "", "Applied", reason
    if any(term in lower for term in (
        "assigned", "search for partners", "reassigned", "submit with partner",
        "appd by ost",
    )):
        return "In progress", "", "Interested", reason
    if "pass" in lower and "partner" in lower:
        return "Referred", "", "Referred", reason
    if any(term in lower for term in ("hold", "undecided", "deferred", "to be discussed")):
        return "On hold", "", "Reviewed", reason
    return text(status) or "Unknown", "", "Reviewed", reason


def read_excel(path: Path) -> list[dict]:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    # Wrike's Excel export can carry a stale ``<dimension ref="A1">`` even
    # though the sheet contains thousands of rows.  In read-only mode
    # openpyxl trusts that metadata and otherwise yields one cell per row.
    # Recalculate the dimensions from the cell stream before reading.
    if sheet.calculate_dimension() == "A1:A1":
        sheet.reset_dimensions()
    rows = sheet.iter_rows(values_only=True)
    try:
        headers = [text(value) for value in next(rows)]
    except StopIteration as exc:
        raise ValueError(f"Historical workbook is empty: {path}") from exc
    required = {"Title", "Item type", "Status", "Start date", "Due date", "Project or folder"}
    missing = sorted(required.difference(headers))
    if missing:
        raise ValueError(f"Historical workbook is missing column(s): {', '.join(missing)}")
    prepared: list[dict] = []
    for values in rows:
        source = dict(zip(headers, values))
        if text(source.get("Item type")).casefold() not in ALLOWED_ITEM_TYPES:
            continue
        title = safe_title(source.get("Title"))
        if not title:
            continue
        status_raw = safe_text(source.get("Status"), limit=128)
        if not status_raw:
            # Rows without an outcome status are containers/meta-tasks rather
            # than usable historical supervision.
            continue
        outcome, won_lost, action, reason = status_mapping(status_raw)
        event_date = iso_date(source.get("Due date")) or iso_date(source.get("Start date"))
        prepared.append({
            "fingerprint": fingerprint_for(title),
            "external_id": "",
            "title": title,
            "source": "OST Analytics",
            "opportunity_url": "",
            "opportunity_type": "",
            "sector": "",
            "verticals": "",
            "brands": "",
            "archetypes": "",
            "geography": "",
            "focus_area": safe_text(source.get("Project or folder"))[:4000],
            "eligibility_text": "",
            "funding": "",
            "company_relevance": "",
            "action_taken": action,
            "status": status_raw or "Unknown",
            "outcome": outcome,
            "won_lost": won_lost,
            "reason": reason,
            "reviewer_comment": "",
            "event_date": event_date,
            "raw_data": {
                "status_group": safe_text(source.get("Status group")),
                "project_or_folder": safe_text(source.get("Project or folder")),
                "start_date": iso_date(source.get("Start date")),
                "due_date": iso_date(source.get("Due date")),
            },
        })
    return prepared


def deduplicate(rows: list[dict]) -> tuple[list[dict], int]:
    selected: dict[str, dict] = {}
    def rank(row: dict) -> tuple[str, int]:
        # Repeated Wrike tasks usually represent the same lead changing state.
        # The latest state is authoritative; decisiveness only breaks a tie.
        outcome_rank = 3 if row["won_lost"] == "won" else 2 if row["won_lost"] == "lost" else 1
        return row.get("event_date") or "", outcome_rank
    for row in rows:
        old = selected.get(row["fingerprint"])
        if old is None or rank(row) > rank(old):
            selected[row["fingerprint"]] = row
    return list(selected.values()), len(rows) - len(selected)


def add_hierarchy(rows: list[dict], batch_size: int = 1000) -> None:
    for start in range(0, len(rows), batch_size):
        batch = rows[start:start + batch_size]
        predictions = classify_hierarchy_batch([
            {"title": row["title"], "summary": row["focus_area"], "source": row["source"]}
            for row in batch
        ])
        if len(predictions) != len(batch):
            raise RuntimeError(
                f"Classifier returned {len(predictions)} prediction(s) for {len(batch)} row(s)"
            )
        for row, prediction in zip(batch, predictions):
            row["opportunity_type"] = prediction.category.value
            row["verticals"] = ", ".join(prediction.verticals)
            row["brands"] = ", ".join(prediction.brands)
            row["archetypes"] = ", ".join(prediction.archetypes)


def _clean_jsonl_row(item: dict, *, line_number: int) -> dict:
    if not isinstance(item, dict):
        raise ValueError(f"JSONL line {line_number} must contain an object")
    title = safe_title(item.get("title"))
    if not title:
        raise ValueError(f"JSONL line {line_number} has no title")
    clean = {field: safe_text(item.get(field)) for field in ROW_FIELDS if field != "event_date"}
    event_date = item.get("event_date")
    if event_date:
        try:
            event_date = date.fromisoformat(str(event_date)).isoformat()
        except ValueError as exc:
            raise ValueError(f"JSONL line {line_number} has an invalid event_date") from exc
    clean["title"] = title
    clean["event_date"] = event_date or None
    clean["fingerprint"] = fingerprint_for(title)
    raw = item.get("raw_data") if isinstance(item.get("raw_data"), dict) else {}
    clean["raw_data"] = {
        key: safe_text(raw.get(key)) if key not in {"start_date", "due_date"} else raw.get(key)
        for key in RAW_DATA_FIELDS
        if raw.get(key) not in (None, "")
    }
    return clean


def read_jsonl(path: Path) -> list[dict]:
    rows: list[dict] = []
    for line_number, line in enumerate(path.read_text(encoding="utf-8").splitlines(), 1):
        if not line.strip():
            continue
        try:
            item = json.loads(line)
        except json.JSONDecodeError as exc:
            raise ValueError(f"Invalid JSON on line {line_number} of {path}") from exc
        rows.append(_clean_jsonl_row(item, line_number=line_number))
    return rows


def write_jsonl(path: Path, rows: list[dict]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="\n") as handle:
        for row in sorted(rows, key=lambda value: (value.get("event_date") or "", value["fingerprint"])):
            handle.write(json.dumps(row, ensure_ascii=False, separators=(",", ":")) + "\n")


def apply_rows(rows: list[dict]) -> dict:
    init_db()
    inserted = updated = unchanged = 0
    with SessionLocal() as db:
        for item in rows:
            row = db.query(HistoricalLead).filter_by(fingerprint=item["fingerprint"]).one_or_none()
            if row is None:
                row = HistoricalLead(fingerprint=item["fingerprint"])
                db.add(row)
                inserted += 1
            else:
                changed = False
                for key in ROW_FIELDS:
                    value = item.get(key, "")
                    if key == "event_date":
                        value = date.fromisoformat(value) if value else None
                    if getattr(row, key) != value:
                        changed = True
                        break
                raw_json = json.dumps(item.get("raw_data", {}), ensure_ascii=False, sort_keys=True)
                if row.raw_data != raw_json:
                    changed = True
                if not changed:
                    unchanged += 1
                    continue
                updated += 1
            for key in ROW_FIELDS:
                value = item.get(key, "")
                if key == "event_date":
                    value = date.fromisoformat(value) if value else None
                setattr(row, key, value)
            row.raw_data = json.dumps(item.get("raw_data", {}), ensure_ascii=False, sort_keys=True)
            row.imported_at = datetime.now(timezone.utc)
        registry = db.get(ModelRegistry, "historical-behaviour-v1") or ModelRegistry(
            version="historical-behaviour-v1", kind="recommendation", status="active"
        )
        registry.metrics = json.dumps({"historical_rows": len(rows), "source": "OST Analytics"})
        db.add(registry)
        db.commit()
    return {
        "inserted": inserted,
        "updated": updated,
        "unchanged": unchanged,
        "total": len(rows),
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("--export-jsonl", type=Path)
    parser.add_argument("--apply", action="store_true")
    parser.add_argument("--skip-classification", action="store_true")
    args = parser.parse_args()
    rows = read_excel(args.source) if args.source.suffix.lower() == ".xlsx" else read_jsonl(args.source)
    source_count = len(rows)
    rows, duplicate_count = deduplicate(rows)
    if not args.skip_classification and any(not row.get("opportunity_type") for row in rows):
        add_hierarchy(rows)
    summary = {
        "source_rows": source_count, "deduplicated_rows": len(rows),
        "duplicates_removed": duplicate_count,
        "won": sum(row["won_lost"] == "won" for row in rows),
        "lost": sum(row["won_lost"] == "lost" for row in rows),
        "undecided": sum(not row["won_lost"] for row in rows),
    }
    if args.export_jsonl:
        write_jsonl(args.export_jsonl, rows)
        summary["exported"] = str(args.export_jsonl)
    if args.apply:
        summary["database"] = apply_rows(rows)
    print(json.dumps(summary, indent=2))


if __name__ == "__main__":
    main()
